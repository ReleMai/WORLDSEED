from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'reference'/'WORLDSEED_Developer_Blueprint_RC7.docx'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd');tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_text(cell,text,bold=False,color=None,size=8.5):
    cell.text='';p=cell.paragraphs[0];r=p.add_run(str(text));r.bold=bold;r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i],h,True,(255,255,255),8.5);shade(t.rows[0].cells[i],'1F4E78')
    # Repeat the header row when Word/LibreOffice carries a table onto a new page.
    trPr=t.rows[0]._tr.get_or_add_trPr()
    tblHeader=OxmlElement('w:tblHeader');tblHeader.set(qn('w:val'),'true');trPr.append(tblHeader)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell_text(cells[i],val,False,None,8.2)
            if ri%2==1:shade(cells[i],'F2F6FA')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):row.cells[i].width=Inches(w)
    return t

def add_bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2');p.add_run(text);return p

def add_num(doc,text):
    p=doc.add_paragraph(style='List Number');p.add_run(text);return p

def add_callout(doc,title,text):
    t=doc.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid';c=t.cell(0,0);shade(c,'EAF2F8')
    c.text='';p=c.paragraphs[0];r=p.add_run(title+'\n');r.bold=True;r.font.color.rgb=RGBColor(31,78,120);r.font.size=Pt(10)
    r=p.add_run(text);r.font.size=Pt(9)

def heading(doc,text,level=1):
    return doc.add_heading(text,level=level)

doc=Document()
sec=doc.sections[0];sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
styles=doc.styles
styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(9.2)
for name,size,color in [('Title',30,(20,55,85)),('Heading 1',18,(20,55,85)),('Heading 2',13,(31,78,120)),('Heading 3',10.5,(55,85,110))]:
    styles[name].font.name='Aptos Display' if name!='Normal' else 'Aptos';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor(*color)
styles['Heading 1'].paragraph_format.space_before=Pt(10);styles['Heading 1'].paragraph_format.space_after=Pt(5)
styles['Heading 2'].paragraph_format.space_before=Pt(8);styles['Heading 2'].paragraph_format.space_after=Pt(3)
styles['Normal'].paragraph_format.space_after=Pt(4)

# Header/footer
for section in doc.sections:
    hp=section.header.paragraphs[0];hp.text='WORLDSEED Developer Blueprint | RC7 Empty Hands';hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(100,100,100)
    fp=section.footer.paragraphs[0];fp.text='WORLDSEED 1.0.0-rc7 | September 2026';fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.runs[0].font.size=Pt(8);fp.runs[0].font.color.rgb=RGBColor(110,110,110)

# Title page
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.space_after=Pt(16)
r=p.add_run('WORLDSEED');r.bold=True;r.font.size=Pt(34);r.font.color.rgb=RGBColor(20,55,85)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Developer-Ready Design & Implementation Blueprint');r.bold=True;r.font.size=Pt(20);r.font.color.rgb=RGBColor(31,78,120)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('RC7 Update - Empty Hands Early-Game Overhaul');r.font.size=Pt(15);r.font.color.rgb=RGBColor(70,90,105)
doc.add_paragraph('')
add_callout(doc,'Defining rule','Model causes and rules; let outcomes emerge. RC7 applies that principle to the opening game: the character begins with nothing, and primitive technology, food, shelter and industry emerge from resources that physically exist in the generated world.')
doc.add_paragraph('')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Status: RC7 implementation and engineering validation complete. Final archive verification is the release gate performed after document generation.\n').bold=True
p.add_run('Engine: 1.0.0-rc7 | Save schema: 3 | Worldgen: 3 | Natural-resource generation: 2\n')
p.add_run('Prepared September 2026')
doc.add_page_break()

heading(doc,'1. Executive Summary')
doc.add_paragraph('WORLDSEED is a terminal-first, persistent fantasy-world simulation in which the world - not a single avatar - is the durable protagonist. Deterministic terrain provides effectively unbounded space; simulation fidelity decreases with distance; only changed or realized state is persisted; ecology, settlements, economies, history, knowledge, religion and magic interact through common events, entities and state.')
doc.add_paragraph('RC7 focuses on a weakness identified through actual play: WORLDSEED was already broad and complex, but its opening hours skipped too much of the physical survival sandbox. RC7 rebuilds that opening around empty hands, material bootstrapping, fieldcraft, useful furniture, hunting/carcass processing, dynamic cooking/nutrition, crop soil roles and richer habitat-aware ecology.')
add_callout(doc,'RC7 release rule','A new character receives no supplies. Fairness comes from a viable landscape, not a starter kit or scripted objective.')
heading(doc,'1.1 Verified RC7 scope',2)
table(doc,['Area','Verified RC7 result'],[
('Starting state','Inventory is empty; point-buy character capability only.'),
('Content','105 items; 52 recipes; 23 structures; 8 crops; 22 natural resources; 17 authored animals; 4 cooking methods.'),
('Regression','156 automated tests passed.'),
('Bootstrap','10/10 independent generated worlds reached fire, timber, carcass processing, cooked food and diet history from zero supplies.'),
('Start balance','20/20 sampled starts were empty and primitive-bootstrap capable; no weak starts.'),
('Performance','~15.11 ms/chunk; ~1.62 ms movement; ~3.97 ms movement + 100x30 render.'),
('Long worlds','3 x 25-year and 1 x 100-year FULL simulations completed with zero invariant errors.'),
], [1.45,5.65])

heading(doc,'2. Product Principles')
principles=[
('Systemic causality','Model resource, need, process and response; do not script the desired story.'),
('Persistent consequences','Player and simulation actions alter the same durable world.'),
('Deterministic potential','Unvisited infinity remains generated potential until required.'),
('Layered fidelity','Individuals locally; cohorts/ledgers/events at distance.'),
('Compositional content','Materials, tool qualities, processes, components and data definitions compose into outcomes.'),
('Discoverability','The engine can know more than the character; UI reveals what the character has learned.'),
('No clutter for realism alone','An item or furniture piece needs a mechanical consumer or meaningful downstream process.'),
('Terminal native','Use information density and navigation hierarchy; do not imitate a graphical UI badly.'),
]
table(doc,['Principle','Implementation meaning'],principles,[1.6,5.5])

heading(doc,'3. Architecture Snapshot')
heading(doc,'3.1 Simulation fidelity',2)
table(doc,['Tier','State','Fidelity'],[
('A','Immediate player area','Exact individuals, movement, combat/anatomy, fluids/fire, local pathing.'),
('B','Nearby','Persistent individuals at reduced decision cadence.'),
('C','Regional','Population cohorts, settlement ledgers, trade, migration, outbreaks.'),
('D','Distant realized','Coarse daily/monthly equations and scheduled events.'),
('E','Unrealized','Deterministic potential only; no persistent population explosion.'),
],[.5,1.55,5.0])
heading(doc,'3.2 Persistent-state rule',2)
doc.add_paragraph('Untouched terrain is reconstructed from seed + generator version + coordinate. Saves primarily store deltas, instantiated entities, discoveries and aggregate living-world state. The same principle applies to RC7 natural resources: deterministic placement is versioned separately, while harvested/regrowing features persist only the changes required to reconstruct their current stage.')
heading(doc,'3.3 Command authority',2)
doc.add_paragraph('The UI requests commands; authoritative simulation validates resources, tools, time, collision and outcomes. Rendering remains observational. RC7 cooking, carcass work, ranged hunting, farming, preservation and construction use the same command/state boundary rather than mutating inventory directly from menus.')

heading(doc,'4. RC7 Zero-Supply Opening')
heading(doc,'4.1 Empty-hands contract',2)
add_bullet(doc,'Inventory is exactly empty at new-character creation.')
add_bullet(doc,'No money, food, tools, medical supplies, seeds, container or hidden origin kit is granted.')
add_bullet(doc,'Character point-buy affects effectiveness, not initial possessions or prescribed story.')
add_bullet(doc,'Spawn scoring can reject unfairly boxed or nonviable starts, but cannot disclose where resources are.')
heading(doc,'4.2 Primitive dependency graph',2)
doc.add_paragraph('The intended first technology graph is:')
add_callout(doc,'Bootstrap','Loose stone + fallen branches + plant fiber/bark -> flakes + cordage -> knife/axe/pick/hammer/digging stick -> fire/shelter/work surface -> timber/mining/hunting/farming -> preserved food/leather/better stations -> advanced production.')
heading(doc,'4.3 Start fairness without handholding',2)
doc.add_paragraph('The final 20-seed start audit found zero weak starts. Within radius 14 the median sampled feature counts were approximately: stone 100.5 cells, branch 36, fiber 39, edible forage 45.5, water 123 and trees 206. The worst nearest observed distance was stone 13, branch 9, fiber 6, edible 5, water 10, tree 3. These numbers are generator diagnostics only; the player is not shown them.')

heading(doc,'5. Item and Craft Economy')
heading(doc,'5.1 Tool-quality model',2)
doc.add_paragraph('Processes request capabilities such as cut, skin, butcher, scrape, chop, mine, hammer, till, dig and sew. A crude or advanced tool can provide the same quality at different strength/speed. This avoids recipe logic that depends on one exact item ID and supports substitution and material progression.')
heading(doc,'5.2 Anti-clutter contract',2)
doc.add_paragraph('An item is accepted only when it is directly usable or has a downstream consumer: process input, construction material, equipment component, storage/fuel/ammunition, agriculture, trade backed by utility, or a knowledge/cultural artifact with consequences. The automated content audit rejects undefined references, missing prose, functionless furniture and obvious orphan content.')
heading(doc,'5.3 Primitive recipe design',2)
table(doc,['Input opportunity','First outputs','What it unlocks'],[
('Loose stone','flakes, hammer/pick heads, sling stones','cutting, mining, striking, ranged ammo'),
('Fallen branches','hafts, fuel, digging stick, spear parts','bound tools, fire, shelter/furniture'),
('Brush/bark/fiber','cordage/bindings','tool assembly, sling, structures'),
('Tree + chopping quality','wood, bark','larger construction, workstations, fuel'),
('Rock/ore + mining quality','stone/ore','masonry, metallurgy'),
],[1.5,2.0,3.6])

heading(doc,'6. Hunting, Combat and Carcass Processing')
doc.add_paragraph('Authored animals carry mass, habitat/climate envelope, health, aggression, group data, named attacks and carcass profiles. Procedural species use the same runtime contracts. Local individuals remain linked to regional population cohorts so local kills affect aggregate ecology.')
heading(doc,'6.1 Carcass lifecycle',2)
table(doc,['Step','Mechanical effect'],[
('Kill','Creates persistent carcass; combat damage contributes to carcass quality loss.'),
('Skin','Produces fresh hide; best performed before butchering.'),
('Butcher','Produces meat, fat and offal; butchering first damages hide yield.'),
('Debone','After butchery, recovers bones and sinew and completes carcass processing.'),
('Delay','Time reduces usable quality; neglected carcasses become unsafe.'),
],[1.0,6.1])
heading(doc,'6.2 Product sinks',2)
table(doc,['Product','Current/future sinks'],[
('Meat','roast, stew, broth, drying, trade'),
('Hide','scrape -> rawhide or tanning -> leather'),
('Bone','tools, needles, broth, fertilizer paths'),
('Sinew','strong cordage/equipment'),
('Fat','cooking, soap, preservation'),
('Offal','food/bait; future feed/fertilizer extensions'),
],[1.0,6.1])

heading(doc,'7. Cooking, Preservation and Nutrition')
heading(doc,'7.1 Ingredient-composed meals',2)
doc.add_paragraph('Current methods are roast, stew, porridge and broth. Methods define station/cookware/water constraints and ingredient roles; the resulting meal derives nutrition from the actual selected ingredients instead of enumerating every combination as a fixed recipe.')
heading(doc,'7.2 Nutrition model',2)
table(doc,['Layer','Tracked values'],[
('Immediate','satiety/hunger and hydration'),
('Meal composition','calories, protein, carbohydrates, fat, fiber'),
('Diet variety','protein, vegetables, grains, fruit, fats'),
('Rolling history','seven-day totals and recent meals'),
('Effects','bounded short-term recovery, stamina or cold-resilience bonuses'),
],[1.3,5.8])
heading(doc,'7.3 Preservation',2)
doc.add_paragraph('Prepared meals have finite shelf life. Carcasses decay. Drying and tanning are timed processing jobs tied to physical stations. Preservation therefore converts perishable opportunity into durable stored value rather than simply creating alternate item names.')

heading(doc,'8. Farming, Forage and Ecology')
heading(doc,'8.1 Crops',2)
table(doc,['Crop set','Distinct data'],[
('Wheat, barley, beans, turnip, cabbage, carrot, onion, flax','growth duration, temperature/water envelope, N/P/K demand, disease family, seed/harvest outputs')
],[2.7,4.4])
doc.add_paragraph('Farming requires a tilling-quality tool. Nutrient use is crop-owned rather than hard-coded, so crop rotation can matter naturally. Wild plant resources can produce seed where appropriate; farming can therefore emerge from exploration rather than a starter packet.')
heading(doc,'8.2 Natural resources',2)
doc.add_paragraph('The 22 authored natural-resource definitions include multiple trees, berries, edible greens/root plants, herbs, fungi, wetland plants, wild grain, clay, branches, loose stone and boulders. Placement uses terrain, climate, season, density, patch scale and deterministic local variation. Structural features such as trees retain identity through stump/sapling/mature stages.')
heading(doc,'8.3 Fauna',2)
doc.add_paragraph('Regional abundance samples actual habitat suitability; local animals materialize on suitable passable cells. The player does not receive exact population statistics without sufficient Naturalism/Survival or reliable reports.')

heading(doc,'9. Functional Construction and Furniture')
table(doc,['Piece','Purpose'],[
('Campfire/hearth','warmth, cooking, light, primitive heat processes'),
('Brush lean-to','shelter/exposure protection'),
('Sleeping mat/bed','rest/recovery'),
('Storage basket/chest','physical storage'),
('Crude work surface/workbench','crafting efficiency'),
('Butcher table','carcass-processing efficiency'),
('Drying rack','hide/meat/herb drying'),
('Tanning frame','leather production'),
('Forge','metallurgy'),
('Walls/gates/roofs/supports','enclosure, protection, structural shelter'),
('Bridge','crossing'),
],[2.0,5.1])
doc.add_paragraph('Purely decorative furniture remains deferred until beauty, prestige, culture or social-space systems can make decoration mechanically meaningful.')

heading(doc,'10. Terminal UX for Complex Crafting')
doc.add_paragraph('The interface is purpose-grouped instead of flat-list driven. Crafting and construction entries show READY/NEEDS before selection. Detail pages use readable material names, station/tool/skill requirements, time, description and the role enabled by the result. Cooking offers a live nutrition preview. Carcass menus show only valid next operations.')
add_bullet(doc,'W/S or arrows navigate; D/Enter selects; A/Esc goes back one layer.')
add_bullet(doc,'The UI describes player-facing consequences, not simulation internals.')
add_bullet(doc,'Long records use pagers; menu selection/scroll state is retained by the UI navigation stack.')

heading(doc,'11. Research Basis')
doc.add_paragraph('RC7 research was used as architectural inspiration rather than copied content.')
table(doc,['Reference','Useful principle adopted'],[
('Cataclysm: Dark Days Ahead JSON documentation','Data-driven items/recipes/requirements and reusable tool qualities instead of exact-tool conditionals.'),
('Vintage Story farming documentation','Distinct soil nutrients and crop rotation as a meaningful agricultural constraint.'),
('Vintage Story meal/nutrition documentation','Ingredient-composed meals and separation of immediate food value from longer-term dietary variety.'),
('UnReal World hideworking documentation/community references','Distinct skinning/butchering/hide-processing stages and quality consequences.'),
('Textual/Ratatui terminal UI patterns','Retained navigation state, explicit screens/modes and model-update-view separation.'),
],[2.3,4.8])
heading(doc,'11.1 Reference URLs',2)
for url in [
'https://docs.cataclysmdda.org/JSON/JSON_INFO.html',
'https://docs.cataclysmdda.org/JSON/ITEM.html',
'https://wiki.vintagestory.at/Farming',
'https://wiki.vintagestory.at/Meal',
'https://wiki.vintagestory.at/Food',
'https://www.unrealworld.fi/wiki/index.php?title=Hideworking',
'https://textual.textualize.io/guide/screens/',
]: add_bullet(doc,url)

heading(doc,'12. Validation and Release Gates')
table(doc,['Gate','Final RC7 status'],[
('Compilation','PASS'),('Automated tests','156/156 PASS'),('Content audit','PASS'),('Zero-item bootstrap','10/10 seeds PASS'),('Start balance','20/20 empty and bootstrap-capable; 0 weak starts'),('Replay','deterministic; RC7 hash f65ea5583575a883f45195a78c6f86d7c5ed64dc'),('Worldgen benchmark','~15.11 ms/chunk'),('Movement benchmark','~1.62 ms/step; ~3.97 ms incl. full render'),('25-year soaks','3/3 seeds; zero invariant errors'),('100-year soak','zero invariant errors; ~1.96 MB state'),
],[2.0,5.1])
heading(doc,'12.1 Invariant focus',2)
for x in ['No negative population/stock/wealth counts.','No orphan settlement/NPC/family/office/infection references.','No unbounded trade-route complete graph.','Unrealized infinity remains deterministic potential.','Content references resolve; items/furniture have utility.','Same seed + same commands remains deterministic.']:add_bullet(doc,x)

heading(doc,'13. Completed vs Remaining')
completed=[
('World/determinism/persistence','Complete foundation'),('Terminal UI/input/navigation','Mature, continuing polish'),('Exploration/discovery','Implemented'),('Primitive early game','RC7 complete'),('Hunting/carcass economy','RC7 complete'),('Dynamic cooking/nutrition','RC7 complete'),('Farming/soil','Substantial, expandable'),('Ecology/procedural fauna','Substantial'),('Construction/rooms/roofs','Substantial'),('Combat/anatomy/equipment','Substantial'),('NPC society/economy/factions','Substantial'),('History/knowledge/religion/magic','Substantial'),('LOD/offline simulation','Implemented/hardened'),('Epidemiology','Implemented')]
table(doc,['System','Status'],completed,[3.0,4.1])
heading(doc,'13.1 Post-RC7 depth priorities',2)
remaining=['Trapping, fishing and physical tracking/sign systems.','Smoking, salting, fermentation, pickling and cold storage.','Livestock/domestication/breeding and deeper animal social/territorial behavior.','Compost/manure, pests, irrigation networks and perennial crops.','Recipe knowledge/proficiencies/experimentation instead of universal recipe knowledge.','Advanced textiles/weaving/clothing construction.','Deeper aesthetics/comfort/prestige so decorative furniture gains real social value.','Search/favorites and keybinding customization for very large craft/build catalogues.','Balance telemetry from real play sessions and additional biome/start sampling.']
for x in remaining:add_bullet(doc,x)

heading(doc,'14. Future Development Rules')
for x in [
'Never give a new character supplies merely to mask an unreachable resource graph.',
'Every new mechanic must interact with at least two existing systems before it is considered a completed feature.',
'Prefer tool qualities, material properties, components and process definitions to one-off item checks.',
'Keep player knowledge separate from simulation knowledge.',
'Keep distant simulation aggregated and event-driven.',
'Preserve deterministic RNG namespace isolation.',
'Run content audit and zero-item bootstrap tests after any primitive resource/recipe change.',
'Run long-world invariants after economy/ecology/AI/history changes.',
'Do not add decorative clutter until another system consumes its aesthetic/social value.',
]:add_num(doc,x)

heading(doc,'Appendix A - RC7 Content Counts')
table(doc,['Content type','Count'],[('Items','105'),('Recipes/process recipes','52'),('Structures/furniture','23'),('Crops','8'),('Natural resource definitions','22'),('Authored animal species','17 + procedural species'),('Dynamic cooking methods','4')],[3.0,4.1])
heading(doc,'Appendix B - Release Commands')
for cmd in [
'python -m compileall -q worldseed',
'python tools/audit_content.py',
'python tools/bootstrap_smoke.py --seeds 10',
'python tools/balance_early_game.py --seeds 20',
'python tools/replay.py tests/fixtures/golden_replay.json',
'python tools/bench_worldgen.py',
'python tools/bench_movement.py',
'python tools/soak_matrix.py --years 25 --seeds 3 --policy full',
'python tools/soak_matrix.py --years 100 --seeds 1 --policy full',
]:
    p=doc.add_paragraph();r=p.add_run(cmd);r.font.name='Consolas';r.font.size=Pt(8.5)

# core properties
props=doc.core_properties;props.title='WORLDSEED Developer Blueprint RC7';props.subject='Living-world sandbox architecture and RC7 early-game implementation status';props.author='WORLDSEED Project';props.keywords='WORLDSEED, sandbox, simulation, terminal, RC7, game design'
doc.save(OUT)
print(OUT)
